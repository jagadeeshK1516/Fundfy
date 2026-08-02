"""Tests for polished document exports — PDF, PPTX, XLSX."""

import os

import pytest

from fundfy.tools.exporters import (
    export_financial_model_xlsx,
    export_markdown,
    export_pdf,
    export_pptx,
    export_pitch_deck_pptx,
    export_business_plan_pdf,
)


SAMPLE_CONTENT = """## Problem
Startups waste months on manual business planning.

## Solution
AI-powered platform that automates business planning and execution.

## Market Size
$50B TAM in business software and consulting.

## Business Model
SaaS subscription: $49/mo starter, $199/mo pro.

## Traction
- 500 beta users
- $10K MRR
- 85% retention rate

## Team
Experienced founders with exits in SaaS and AI.

## The Ask
Raising $2M seed round to scale go-to-market.
"""


class TestExportPDF:
    """Tests for PDF export."""

    def test_export_pdf_creates_file(self):
        """Test that PDF export creates a valid file."""
        filepath = export_pdf("Test Business Plan", SAMPLE_CONTENT)
        assert os.path.exists(filepath)
        assert filepath.endswith(".pdf")
        assert os.path.getsize(filepath) > 0
        os.unlink(filepath)

    def test_export_business_plan_pdf(self):
        """Test business plan PDF export."""
        filepath = export_business_plan_pdf("My Business Plan", SAMPLE_CONTENT)
        assert os.path.exists(filepath)
        assert filepath.endswith(".pdf")
        assert os.path.getsize(filepath) > 1000  # Should be a substantial file
        os.unlink(filepath)


class TestExportPPTX:
    """Tests for PPTX export."""

    def test_export_pptx_creates_file(self):
        """Test that PPTX export creates a valid file."""
        filepath = export_pptx("Pitch Deck", SAMPLE_CONTENT)
        assert os.path.exists(filepath)
        assert filepath.endswith(".pptx")
        assert os.path.getsize(filepath) > 0
        os.unlink(filepath)

    def test_export_pitch_deck_pptx(self):
        """Test pitch deck PPTX with branding."""
        filepath = export_pitch_deck_pptx("Fundfy Pitch", SAMPLE_CONTENT)
        assert os.path.exists(filepath)
        assert filepath.endswith(".pptx")
        assert os.path.getsize(filepath) > 1000
        os.unlink(filepath)

    def test_pptx_has_multiple_slides(self):
        """Test that PPTX has multiple slides (one per section)."""
        from pptx import Presentation

        filepath = export_pptx("Multi Slide", SAMPLE_CONTENT)
        prs = Presentation(filepath)
        # Title + sections + Thank You
        assert len(prs.slides) >= 5
        os.unlink(filepath)


class TestExportXLSX:
    """Tests for XLSX financial model export."""

    def test_export_xlsx_creates_file(self):
        """Test that XLSX export creates a valid file."""
        filepath = export_financial_model_xlsx("Financial Model", {})
        assert os.path.exists(filepath)
        assert filepath.endswith(".xlsx")
        assert os.path.getsize(filepath) > 0
        os.unlink(filepath)

    def test_xlsx_has_multiple_sheets(self):
        """Test that XLSX has the required sheets."""
        from openpyxl import load_workbook

        filepath = export_financial_model_xlsx("Model", {})
        wb = load_workbook(filepath)
        sheet_names = wb.sheetnames
        assert "Summary" in sheet_names
        assert "Revenue Projections" in sheet_names
        assert "Cost Structure" in sheet_names
        assert "P&L" in sheet_names
        assert "Cash Flow" in sheet_names
        assert "Assumptions" in sheet_names
        os.unlink(filepath)

    def test_xlsx_with_string_data(self):
        """Test XLSX export when data is a string."""
        filepath = export_financial_model_xlsx("Model", "revenue data here")
        assert os.path.exists(filepath)
        os.unlink(filepath)

    def test_xlsx_with_custom_data(self):
        """Test XLSX export with custom data dict."""
        data = {
            "summary": {"Total Revenue": 1000000},
            "months": list(range(1, 7)),
            "revenue": [50000, 60000, 70000, 80000, 90000, 100000],
            "costs": {"Engineering": 150000, "Marketing": 50000},
            "pl": {
                "Revenue": [500000, 1000000, 2000000],
                "Net Income": [-50000, 200000, 500000],
            },
            "cash_flow": [-20000, -10000, 0, 10000, 20000, 30000],
            "assumptions": {"Growth Rate": "20%"},
        }
        filepath = export_financial_model_xlsx("Custom Model", data)
        assert os.path.exists(filepath)
        assert os.path.getsize(filepath) > 1000
        os.unlink(filepath)


class TestExportMarkdown:
    """Tests for Markdown export."""

    def test_export_markdown_creates_file(self):
        """Test that Markdown export creates a valid file."""
        filepath = export_markdown("Test Doc", SAMPLE_CONTENT)
        assert os.path.exists(filepath)
        assert filepath.endswith(".md")
        with open(filepath, "r") as f:
            content = f.read()
        assert "# Test Doc" in content
        os.unlink(filepath)
