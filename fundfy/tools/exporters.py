"""Document export functions — PDF, DOCX, PPTX, Markdown."""

import os
import uuid

from fundfy.config import settings


def _ensure_dir() -> str:
    """Ensure the generated files directory exists and return its path."""
    dir_path = settings.generated_files_dir
    os.makedirs(dir_path, exist_ok=True)
    return dir_path


def export_markdown(title: str, content: str) -> str:
    """Export content as a Markdown file. Returns the file path."""
    dir_path = _ensure_dir()
    filename = f"{uuid.uuid4()}.md"
    filepath = os.path.join(dir_path, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(f"# {title}\n\n{content}")
    return filepath


def export_pdf(title: str, content: str) -> str:
    """Export content as a PDF file using reportlab. Returns the file path."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer

    dir_path = _ensure_dir()
    filename = f"{uuid.uuid4()}.pdf"
    filepath = os.path.join(dir_path, filename)

    doc = SimpleDocTemplate(filepath, pagesize=letter)
    styles = getSampleStyleSheet()
    story = []

    # Title
    story.append(Paragraph(title, styles["Title"]))
    story.append(Spacer(1, 12))

    # Content — split by newlines into paragraphs
    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            # Escape XML special chars for reportlab
            paragraph = paragraph.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            story.append(Paragraph(paragraph, styles["Normal"]))
            story.append(Spacer(1, 6))

    doc.build(story)
    return filepath


def export_docx(title: str, content: str) -> str:
    """Export content as a DOCX file using python-docx. Returns the file path."""
    from docx import Document

    dir_path = _ensure_dir()
    filename = f"{uuid.uuid4()}.docx"
    filepath = os.path.join(dir_path, filename)

    doc = Document()
    doc.add_heading(title, 0)

    for paragraph in content.split("\n"):
        paragraph = paragraph.strip()
        if paragraph:
            if paragraph.startswith("## "):
                doc.add_heading(paragraph[3:], level=2)
            elif paragraph.startswith("# "):
                doc.add_heading(paragraph[2:], level=1)
            else:
                doc.add_paragraph(paragraph)

    doc.save(filepath)
    return filepath


def export_pptx(title: str, content: str) -> str:
    """Export content as a PPTX file using python-pptx. Splits by ## headings into slides."""
    from pptx import Presentation
    from pptx.util import Inches

    dir_path = _ensure_dir()
    filename = f"{uuid.uuid4()}.pptx"
    filepath = os.path.join(dir_path, filename)

    prs = Presentation()

    # Title slide
    slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(slide_layout)
    slide.shapes.title.text = title

    # Split content by ## headings into slides
    sections = content.split("## ")
    for section in sections:
        section = section.strip()
        if not section:
            continue

        lines = section.split("\n", 1)
        slide_title = lines[0].strip()
        slide_body = lines[1].strip() if len(lines) > 1 else ""

        slide_layout = prs.slide_layouts[1]  # Title + Content
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = slide_title
        if slide_body and slide.placeholders[1]:
            slide.placeholders[1].text = slide_body

    prs.save(filepath)
    return filepath
